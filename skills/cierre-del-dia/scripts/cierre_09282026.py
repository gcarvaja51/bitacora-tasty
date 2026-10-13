# Informe cierre del día 28-sep-2026: .docx propio + log (resultado, llamada_siguiente) + Excel.
import json, os, shutil
from datetime import datetime
from zoneinfo import ZoneInfo
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

BASE = os.path.join(r"C:\Users\gcarv\Documents\CARPETA PERSONAL", "01. guillermo carvajal", "01_Sigma",
                    "mentoria alejandro", "premercados alejandro")
DOC = os.path.join(BASE, "documentos premercado", "09282026_cierre del dia.docx")
LOG = os.path.join(BASE, "control premercado", "premercado_hipotesis_log.json")
XLS = os.path.join(BASE, "control premercado", "control_premercado.xlsx")
assert not os.path.exists(DOC), "el informe de hoy ya existe: no se sobrescribe"
HORA = datetime.now(ZoneInfo("America/New_York")).strftime("%H:%M")

doc = Document()
for sec in doc.sections:
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(0.75)
for sn in ("Normal", "Heading 1", "Heading 2", "Heading 3", "List Bullet"):
    st = doc.styles[sn]; st.font.name = "Tahoma"; st.font.size = Pt(11)
    rpr = st.element.get_or_add_rPr(); rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = rpr.makeelement(qn('w:rFonts'), {}); rpr.append(rf)
    rf.set(qn('w:eastAsia'), 'Tahoma')
    if sn in ("Normal", "List Bullet"):
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

def fuente(run, size=11, bold=None):
    run.font.name = "Tahoma"; run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = rpr.makeelement(qn('w:rFonts'), {}); rpr.append(rf)
    rf.set(qn('w:eastAsia'), 'Tahoma')
    if bold is not None: run.bold = bold

def H(t, l=2):
    h = doc.add_heading("", level=l); fuente(h.add_run(t))
def P(t):
    p = doc.add_paragraph(); fuente(p.add_run(t)); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
def B(t):
    p = doc.add_paragraph(style="List Bullet"); fuente(p.add_run(t)); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
def tabla(cols, filas):
    t = doc.add_table(rows=1, cols=len(cols)); t.style = "Light Grid Accent 1"
    for i, c in enumerate(cols):
        cell = t.rows[0].cells[i]; cell.text = ""
        fuente(cell.paragraphs[0].add_run(c), 9.5, True); cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    for f in filas:
        row = t.add_row()
        for i, v in enumerate(f):
            cell = row.cells[i]; cell.text = ""
            fuente(cell.paragraphs[0].add_run(str(v)), 9.5); cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.add_paragraph()

H("Cierre del día — lunes 28 de septiembre de 2026", 1)
P(f"Hora del informe: {HORA} ET. Día SIN TRADE para las estadísticas del premercado (D19): declaraciones de Trump "
  "a las 14:00 ET, evento que Investing califica con 3 estrellas. Al cierre, la pestaña Hoy del calendario (filtro "
  "EE.UU., 3 estrellas) seguía sin listarlo y no mostraba ningún otro dato: no apareció nada nuevo.")

H("Qué hizo el mercado", 3)
P("Abrió en 7.722 y cerró en 7.684: −59 puntos (−0,77 %) contra el cierre del viernes (7.743). Máximo 7.724 a las "
  "12:15, mínimo 7.667 a las 10:45. Rango de 58 puntos. El VIX casi no se movió: 16,2 a la apertura, 16,1 al cierre, "
  "máximo 16,6. Caída sin miedo.")
B("09:30 — Abrió con un gap de −22 y la primera vela bajó hasta 7.698, cierre 7.705: fue directo al Put Wall "
  "(7.700) del premercado.")
B("09:45 — Rebote a 7.712. Con ese cierre el cruce EMA10/EMA20 de 15m se dio vuelta a la baja.")
B("10:00 — La tercera vela cerró en 7.702, dentro del corredor Neutral (7.700-7.735), apoyada en el 7.700.")
B("10:30 — Perdió el 7.700: cierre de 15m en 7.682, bajo 7.697 (umbral con el colchón). Activó el Bajista y tocó "
  "su T1 (7.685) en la misma vela, con mínimo 7.682.")
B("10:45 — Mínimo del día, 7.667: T2 del Bajista (7.675) cumplido. Cierre 7.671.")
B("11:00 a 12:00 — Una hora de rango entre 7.671 y 7.687, sobre la EMA20 diaria (7.681).")
B("12:15 — Vela de +33 hasta 7.724, el máximo del día, sin un dato de 3 estrellas detrás. Se devolvió entera en "
  "la vela siguiente (7.700).")
B("13:00 — Segundo intento a 7.721 y cierre en 7.718, 5 puntos bajo la invalidación del Bajista (7.723 con el "
  "colchón). A las 13:15 ya estaba otra vez en 7.700.")
B("13:30 a 14:45 — Pin de hora y cuarto entre 7.693 y 7.705, clavado en el 7.700.")
B("14:00 — Declaraciones de Trump (3 estrellas, sin pronóstico): la vela de las 14:00 tuvo un rango de 6 puntos. "
  "Sin reacción.")
B("15:00 a 16:00 — Perdió el 7.700 por última vez y derivó a 7.678. Cierre en 7.684, pegado a los muros de Sigma "
  "del cierre (7.680-7.685).")

H("Escenarios del premercado contra lo que pasó", 3)
tabla(["Escenario", "Prob.", "¿Se validó?", "Nota"],
 [["Neutral — corredor 7.700-7.735", "37 %", "NO",
   "Aguantó la primera hora apoyado en el 7.700 y se rompió por abajo al cierre de las 10:30 (7.682)."],
  ["Bajista — pierde el Put Wall", "32 %", "SÍ, completo",
   "Activó al cierre de las 10:30 (7.682, bajo 7.697 con el colchón). T1 7.685 en la misma vela y T2 7.675 a las "
   "10:45 (mínimo 7.667). No se invalidó: el mejor cierre posterior fue 7.718 a las 13:00, bajo 7.723."],
  ["Alcista — recupera el Flip", "31 %", "NO",
   "Nunca activó: el cierre de 15m más alto del día fue 7.718, lejos de 7.738."]])

H("¿Acertamos?", 3)
P("NO. El favorito era el Neutral (37 %) y se validó el Bajista (32 %), que cumplió los dos objetivos en la primera "
  "media hora después de activarse. La diferencia entre los tres escenarios era de 6 puntos porcentuales: el "
  "premercado no tenía una lectura clara. Como el día es D19 (Trump a las 14:00), no suma a la asertividad ni a los "
  "totales. La llamada del viernes para hoy (ALCISTA sobre 7.755, invalidada bajo 7.730) falló: se invalidó con el "
  "gap de la apertura, sin llegar a activarse.")

H("Cómo quedó Sigma al cierre", 3)
P("Lectura de las 16:02 ET (la más reciente de status.json; la de ~16:05 todavía no estaba), cadena del 28-sep, que "
  "venció hoy. El régimen se dio vuelta: el GEX pasó de +6,7 B en el premercado a −33,7 B, y el DEX de +8,5 B a "
  "−13,2 B. Los muros se juntaron sobre el cierre (Put Wall 7.680, Call Wall y MVS 7.685) y el Flip quedó en 7.753, "
  "70 puntos arriba; Max Pain 7.725. Esa cadena ya venció, así que los muros de mañana se leen en el premercado. Lo "
  "que sirve es el tono: el día terminó en gamma negativo, con los dealers acompañando los movimientos en vez de "
  "frenarlos.")

H("Qué puede pasar mañana, martes 29 de septiembre", 3)
P("Estructura al cierre:")
B("Diario: alcista pero en el límite. El cierre (7.684) quedó bajo la EMA10 (7.694) y 3 puntos sobre la EMA20 "
  "(7.681); la EMA10 sigue arriba de la EMA20 y la EMA50 está en 7.631. El mínimo de hoy (7.667) repitió la zona "
  "del mínimo del jueves 24 (7.663).")
B("30 minutos: debajo de la EMA10 (7.697), la EMA20 (7.704) y la EMA50 (7.708), y apenas bajo la EMA200 (7.687). El "
  "cruce EMA10/EMA20 de 30m se dio vuelta a la baja en la vela de las 10:30 de hoy.")
B("15 minutos (el indicador principal): el cruce EMA10/EMA20 es bajista desde la vela de las 09:45 de hoy. EMA10 "
  "7.692 y EMA20 7.697. La distancia llegó a −12,7 a las 11:30, se cerró hasta −2,0 a las 14:45 sin llegar a "
  "cruzarse, y en la última hora se volvió a abrir a −4,7. Para darlo vuelta al alza hace falta un cierre de 15m "
  "sobre 7.697 y que se sostenga.")
B("Fuera de hora: el ES de diciembre estaba en 7.740 a las 15:55 ET. Con la base de hoy (60,7, mediana de 26 velas "
  "de 15m pareadas) equivale a 7.680: apertura implícita provisoria de −4, plana. Falta toda la noche.")
P("Niveles para mañana (entre paréntesis, en tu pantalla VANTAGE:SP500, con +10):")
tabla(["Nivel", "Qué es"],
 [["7.725 (7.735)", "Máximo de hoy (7.724) y Max Pain de Sigma"],
  ["7.715 (7.725)", "Memoria fuerte (8 giros, el último el 24-sep) y techo de los dos rebotes de hoy (7.718-7.721)"],
  ["7.705 (7.715)", "EMA20 y EMA50 de 30m (7.704 / 7.708) y borde de arriba del pin de la tarde"],
  ["7.700 (7.710)", "El pivote de hoy: soporte de la mañana y pin de 13:30 a 14:45"],
  ["7.695 (7.705)", "EMA10 diaria (7.694) y EMA20 de 15m (7.697)"],
  ["7.685 (7.695)", "Cierre, EMA200 de 30m (7.687) y muros de Sigma al cierre (7.680-7.685)"],
  ["7.680 (7.690)", "EMA20 diaria (7.681): el soporte de la tendencia de fondo"],
  ["7.665 (7.675)", "Mínimo de hoy (7.667) y del jueves 24 (7.663); memoria fuerte en 7.660 (10 giros)"],
  ["7.650 (7.660)", "Cierre del viernes 18 (7.650) y memoria media (7 giros)"],
  ["7.640 (7.650)", "Memoria media (giros del 24-ago y 9-sep); la EMA50 diaria está debajo, en 7.631"]])
P("Catalizadores de 3 estrellas (Investing, filtro EE.UU., pestaña Mañana): confianza del consumidor de The "
  "Conference Board (pronóstico 90,1, previo 89,4) y encuesta JOLTS de ofertas de empleo (pronóstico 7,230M, previo "
  "7,271M), los dos a las 10:00 ET. Ninguno a las 10:15 o después, así que mañana SÍ se opera. Los dos caen dentro "
  "de la tercera vela: la hipótesis activa se califica con el dato recién publicado.")
P("Llamada para mañana: BAJISTA, convicción 55 %. Se activa con un cierre de 15m bajo 7.665 (7.675 en pantalla): "
  "pierde el mínimo de hoy y el del jueves, en gamma negativo y con los cruces de 15m y 30m a la baja. T1 7.650 "
  "(cierre del 18-sep), T2 7.640 (memoria media, con la EMA50 diaria debajo). Se invalida con un cierre de 15m sobre "
  "7.705 (7.715): recupera el 7.700 y las EMA20/50 de 30m, y el cruce de 15m se daría vuelta. El lado alcista tiene "
  "argumentos reales: la tendencia diaria sigue intacta (EMA10 sobre EMA20), el precio cerró encima de la EMA20 "
  "diaria, la zona 7.660-7.665 es un piso probado dos veces en cuatro días, y hoy hubo dos rebotes de +30 en la "
  "tarde. Además los datos de las 10:00 pueden mover el precio para cualquier lado. Por eso la convicción es baja. "
  "Esta llamada no se edita: mañana se puntúa sola.")
doc.save(DOC)

shutil.copy(LOG, LOG + ".bak_20260928_cierre")
L = json.load(open(LOG, encoding='utf-8'))
e = [x for x in L if x['fecha'] == '2026-09-28'][0]
e['resultado'] = {
    "escenario_validado": "bajista (activo al cierre 15m de las 10:30 en 7.682,22; T1 7.685 en la misma vela con min 7.681,64; T2 7.675 a las 10:45 con min 7.666,60; no invalidado)",
    "apertura_real": 7721.70, "maximo": 7724.15, "minimo": 7666.60, "cierre_real": 7683.92,
    "acierto": "no",
    "acierto_nota": "Favorito neutral 37% -> NO; se valido el bajista 32%. Dia D19 (Trump 14:00 ET, 3 estrellas): fuera de asertividad y totales. La llamada del 25-sep para hoy (ALCISTA sobre 7.755) se invalido con la apertura (cierre 09:30 = 7.704,94 < 7.730) sin activarse.",
    "veredicto_mecanico": "Bajista VALIDADO: cierre 15m 10:30 = 7.682,22 < 7.697 (7.700-3); min 7.681,64 <= T1 7.685; min 10:45 7.666,60 <= T2 7.675; max cierre posterior 7.717,53 (13:00) < 7.723 (7.720+3). Alcista NO_ACTIVO (max cierre 7.717,53 < 7.738). Neutral INVALIDADO al cierre de las 10:30.",
    "leccion_aprendida": "Con tres escenarios separados por 6 puntos porcentuales el favorito no tenia ventaja real. El gap de -22 dejo la apertura en el Put Wall y el soporte aguanto solo una hora. GEX paso de +6,7B (premercado) a -33,7B al cierre: otra vez el regimen de las 08:30 no anticipo el dia. Las declaraciones de Trump (14:00) no movieron el precio.",
    "catalizadores_resultado": [
        {"evento": "Declaraciones de Trump", "actual": None, "pronostico": None, "reaccion": "vela 14:00 rango 6 pts (7.698,74-7.704,27), sin reaccion"},
    ],
    "sigma_cierre": {"capturado": "2026-09-28T20:02:25Z", "call_wall": 7685, "put_wall": 7680, "gamma_flip": 7753, "mvs": 7685, "max_pain": 7725, "net_gex_b": -33.7, "net_dex_b": -13.22, "regimen": "NEGATIVO"},
}
e['llamada_siguiente'] = {
    "para_fecha": "2026-09-29", "emitida_et": "2026-09-28T" + HORA, "spot_cierre": 7683.92,
    "direccion": "BAJISTA", "conviccion": 55,
    "activa": {"tipo": "cierre_15m_bajo", "nivel": 7665},
    "invalida": {"tipo": "cierre_15m_sobre", "nivel": 7705},
    "t1": 7650, "t2": 7640,
    "tesis": "Gamma negativo al cierre (-33,7B), cruce EMA10/20 bajista en 15m (desde 09:45) y 30m (desde 10:30), cierre bajo EMA10 diaria y EMAs de 30m. Si pierde el minimo 7.667 y el del 24-sep (7.663) busca el cierre del 18-sep (7.650) y la memoria de 7.640. Conviccion baja: tendencia diaria intacta, cierre sobre la EMA20 diaria, piso 7.660-7.665 probado y datos 3 estrellas a las 10:00 (Conference Board, JOLTS).",
    "resultado": None,
}
json.dump(L, open(LOG, 'w', encoding='utf-8'), ensure_ascii=False, indent=2)

import openpyxl
shutil.copy(XLS, XLS.replace(".xlsx", "_bak_20260928_cierre.xlsx"))
wb = openpyxl.load_workbook(XLS); ws = wb['Control Premercado']
assert ws.cell(69, 1).value == '2026-09-28'
ws.cell(69, 4).value = (ws.cell(69, 4).value or '') + (f' || CIERRE DEL DIA ({HORA} ET): Bajista VALIDADO completo (activo 10:30 en 7.682; T1 misma vela, T2 10:45; min 7.667). '
    'Cierre 7.684 (-59, -0,77%). Acierto NO (favorito neutral 37%). Dia D19, no suma. Sigma cierre GEX -33,7B, muros 7.680-7.685. '
    'Llamada 29-sep: BAJISTA 55% bajo 7.665, invalida sobre 7.705, T1 7.650 T2 7.640. Manana 3 estrellas: Conference Board y JOLTS 10:00 (se opera).')
wb.save(XLS)
print('OK', DOC, HORA)
