# Corrige el informe de hoy con la lectura de Sigma de las 16:05 ET (llegó después de generarlo).
import json, os
from docx import Document

BASE = os.path.join(r"C:\Users\gcarv\Documents\CARPETA PERSONAL", "01. guillermo carvajal", "01_Sigma",
                    "mentoria alejandro", "premercados alejandro")
DOC = os.path.join(BASE, "documentos premercado", "09282026_cierre del dia.docx")
LOG = os.path.join(BASE, "control premercado", "premercado_hipotesis_log.json")
XLS = os.path.join(BASE, "control premercado", "control_premercado.xlsx")

NUEVO = ("Lectura de las 16:05 ET, cadena del 28-sep, que venció hoy. El régimen se dio vuelta: el GEX pasó de +6,7 B "
  "en el premercado a −27,7 B, y el DEX de +8,5 B a −13,9 B. Los muros se juntaron sobre el cierre (Put Wall 7.680, "
  "Call Wall y MVS 7.685) y el Flip quedó en 7.744, 60 puntos arriba; Max Pain 7.725. Esa cadena ya venció, así que "
  "los muros de mañana se leen en el premercado. Lo que sirve es el tono: el día terminó en gamma negativo, con los "
  "dealers acompañando los movimientos en vez de frenarlos.")
doc = Document(DOC); n = 0
for p in doc.paragraphs:
    if p.text.startswith("Lectura de las 16:02 ET"):
        for r in p.runs[1:]: r.text = ""
        p.runs[0].text = NUEVO; n += 1
assert n == 1
doc.save(DOC)

L = json.load(open(LOG, encoding='utf-8'))
e = [x for x in L if x['fecha'] == '2026-09-28'][0]
e['resultado']['sigma_cierre'] = {"capturado": "2026-09-28T20:05:59Z", "call_wall": 7685, "put_wall": 7680, "gamma_flip": 7744, "mvs": 7685, "max_pain": 7725, "net_gex_b": -27.66, "net_dex_b": -13.87, "regimen": "NEGATIVO"}
e['resultado']['leccion_aprendida'] = e['resultado']['leccion_aprendida'].replace("-33,7B", "-27,7B")
e['llamada_siguiente']['tesis'] = e['llamada_siguiente']['tesis'].replace("-33,7B", "-27,7B")
json.dump(L, open(LOG, 'w', encoding='utf-8'), ensure_ascii=False, indent=2)

import openpyxl
wb = openpyxl.load_workbook(XLS); ws = wb['Control Premercado']
assert ws.cell(69, 1).value == '2026-09-28'
ws.cell(69, 4).value = ws.cell(69, 4).value.replace("GEX -33,7B", "GEX -27,7B")
wb.save(XLS)
print('OK')
